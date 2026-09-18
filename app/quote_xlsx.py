"""Builds the quotation .xlsx exactly like the AppSheet template.

The layout is a faithful rebuild of the sample workbook (see
data/SAMPLE_QUOTE_REFERENCE.xlsx): same fonts, column widths, merges, fills,
number formats and footer text.  The only thing that is dynamic is the item
block - it grows in 3-row steps so a quotation can carry any number of lines,
while still rendering the original 7 blank slots on a short quote.
"""

import os
import re
import subprocess
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- constants lifted from the template ------------------------------------

HEADER_ROW = 16          # No. | Description | GST | Qty | Rate | Amount
FIRST_ITEM_ROW = 17      # each item occupies 3 rows: value, model, spacer
ROWS_PER_ITEM = 3
MIN_ITEM_SLOTS = 7       # keeps a short quote looking like the original page

CURRENCY_FMT = '[$₹]#,##0.00'
RATE_FMT = '#,##0.00'
PERCENT_FMT = '0.00%'
DATE_FMT = 'M/d/yyyy H:mm:ss'

THIN = Side(style="thin")
WHITE = "FFFFFFFF"
BLACK_FILL = PatternFill("solid", fgColor="FF000000")
RED_FILL = PatternFill("solid", fgColor="FFFF0000")
BLUE_FILL = PatternFill("solid", fgColor="FF0000FF")

COL_WIDTHS = {"A": 5.38, "B": 46.25, "C": 9.13, "D": 5.75, "E": 10.38,
              "F": 16.75, "G": 8.0, "H": 10.13, "I": 8.0}

FOOTER_NOTES = [
    "1) Goods Once sold will not be taken back",
    "2) Our responsibility ceases once the goods leave our godown",
    "3) Payment by crossed cheque is requested",
    "4) Subject to Mumbai jurisdiction (s) E. & O.E",
]

GST_DECLARATION = [
    "I / We hereby certify that my / our registration certificate under the GST",
    " is in force  on  the  date on which the sale of the goods specified in  this  tax  ",
    "invoice  is made by me / us and  that  the  transaction  of  the sale covered",
    "by this  tax  invoice  has been effected by  me / us  and  it shall be accounted ",
    "for in the turnover of sales while filing of return and the due tax,  if any, ",
    "payable on the sale has been paid or shall be paid",
]


def _cambria(size=10, bold=False, color=None):
    return Font(name="Cambria", sz=size, b=bold, color=color)


def _box(cell, left=False, right=False, top=False, bottom=False):
    cell.border = Border(
        left=THIN if left else None,
        right=THIN if right else None,
        top=THIN if top else None,
        bottom=THIN if bottom else None,
    )


def _merge_write(ws, ref, value, font=None, align=None, fill=None, fmt=None):
    """Merge `ref` and write `value` into its top-left cell."""
    ws.merge_cells(ref)
    cell = ws[ref.split(":")[0]]
    cell.value = value
    if font:
        cell.font = font
    if align:
        cell.alignment = align
    if fill:
        cell.fill = fill
        # A fill on a merged range only paints the anchor cell unless every
        # cell in the range carries it too.
        for row in ws[ref]:
            for c in row:
                c.fill = fill
    if fmt:
        cell.number_format = fmt
    return cell


def safe_filename(text, fallback="QUOTATION"):
    text = re.sub(r"[\\/:*?\"<>|\r\n\t]+", " ", (text or "")).strip()
    text = re.sub(r"\s+", " ", text)
    return (text or fallback)[:120]


def build_filename(quote):
    """`PARTY NAME FIRST ITEM DESCRIPTION.xlsx`, same as the AppSheet output."""
    party = quote.get("party_name") or "QUOTATION"
    items = [i for i in quote.get("items", []) if (i.get("description") or i.get("model"))]
    suffix = items[0].get("description") or items[0].get("model") if items else ""
    return safe_filename(f"{party} {suffix}".strip()) + ".xlsx"


def generate_quote_xlsx(quote, out_path):
    """Write the quotation workbook.

    `quote` is a dict with: company{name,address,tel,bank_details},
    party_name, party_address, quote_date (datetime), salesperson,
    salesperson_phone, order_status, dispatch, advance, gst_percent, terms,
    items[{model, description, gst_percent, qty, rate, amount}],
    subtotal, gst_amount, grand_total, balance.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    company = quote.get("company") or {}
    items = [i for i in quote.get("items", []) if (i.get("description") or i.get("model"))]
    slots = max(len(items), MIN_ITEM_SLOTS)

    # ---- title ------------------------------------------------------------
    _merge_write(ws, "A1:F1", "QUOTATION",
                 font=Font(name="Times New Roman", sz=12, b=True, color=WHITE),
                 align=Alignment(horizontal="center"), fill=BLACK_FILL)

    # ---- date / prepared by ----------------------------------------------
    ws["A2"] = "DATE:-"
    ws["A2"].font = _cambria()
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["B2"] = quote.get("quote_date")
    ws["B2"].font = _cambria(bold=True)
    ws["B2"].number_format = DATE_FMT
    ws["B2"].alignment = Alignment(horizontal="left")
    ws["C2"] = "Prepared by"
    ws["C2"].font = _cambria()
    prepared = " ".join(x for x in [quote.get("salesperson"),
                                    quote.get("salesperson_phone")] if x)
    ws["E2"] = prepared
    ws["E2"].font = _cambria()

    # ---- our company ------------------------------------------------------
    _merge_write(ws, "A3:F3", company.get("name", ""),
                 font=Font(name="Times New Roman", sz=24, b=True),
                 align=Alignment(horizontal="center", vertical="center"))
    ws.row_dimensions[3].height = 30

    company_block = "\n".join(
        x for x in [company.get("address", ""), company.get("address2", "")] if x
    )
    _merge_write(ws, "A4:F6", company_block,
                 font=Font(name="Arial", sz=8),
                 align=Alignment(horizontal="center", vertical="center", wrap_text=True))
    for r in (4, 5, 6):
        ws.row_dimensions[r].height = 12.75

    # ---- party ------------------------------------------------------------
    ws["A7"] = "M/S:"
    ws["A7"].font = _cambria(bold=True)
    _merge_write(ws, "B7:F7", quote.get("party_name", ""),
                 font=_cambria(14),
                 align=Alignment(horizontal="center", vertical="center"))

    ws["A8"] = "ADD:"
    ws["A8"].font = _cambria(bold=True)
    _merge_write(ws, "B8:F11", quote.get("party_address", ""),
                 font=_cambria(9),
                 align=Alignment(horizontal="center", vertical="center", wrap_text=True))

    # ---- status block -----------------------------------------------------
    status_rows = [
        ("ORDER STATUS:-", quote.get("order_status", ""), None),
        ("ADVANCE:-", quote.get("advance") or None, CURRENCY_FMT),
        ("BALANCE PAYABLE:-", quote.get("balance") or None, CURRENCY_FMT),
        ("DISPATCH:-", quote.get("dispatch", ""), None),
    ]
    for offset, (label, value, fmt) in enumerate(status_rows):
        row = 12 + offset
        _merge_write(ws, f"A{row}:B{row}", label, font=_cambria(bold=True),
                     align=Alignment(horizontal="right"))
        cell = _merge_write(ws, f"C{row}:F{row}", value, font=_cambria(bold=True),
                            align=Alignment(horizontal="left"))
        if fmt and value is not None:
            cell.number_format = fmt
        ws.row_dimensions[row].height = 12.75

    # ---- item table header ------------------------------------------------
    headers = ["No.", "Description", "GST", "Qty", "Rate", "Amount "]
    for idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=idx, value=text)
        cell.font = _cambria(bold=True)
        cell.alignment = Alignment(horizontal="center")
        _box(cell, left=True, right=(idx == 6), top=True, bottom=True)
    ws.row_dimensions[HEADER_ROW].height = 14.25

    # ---- items ------------------------------------------------------------
    last_item_row = FIRST_ITEM_ROW + ROWS_PER_ITEM * slots - 1
    for row in range(FIRST_ITEM_ROW, last_item_row + 1):
        for col in range(1, 7):
            _box(ws.cell(row=row, column=col), left=True, right=(col == 6))
        ws.row_dimensions[row].height = 14.25 if row < FIRST_ITEM_ROW + 2 else 12.75

    for index, item in enumerate(items):
        row = FIRST_ITEM_ROW + ROWS_PER_ITEM * index
        qty = float(item.get("qty") or 0)
        rate = float(item.get("rate") or 0)
        amount = float(item.get("amount") if item.get("amount") is not None else qty * rate)

        serial = ws.cell(row=row, column=1, value=index + 1)
        serial.font = _cambria()
        serial.alignment = Alignment(horizontal="center")

        desc = ws.cell(row=row, column=2, value=item.get("description", ""))
        desc.font = _cambria()

        gst = ws.cell(row=row, column=3, value=float(item.get("gst_percent") or 0) / 100.0)
        gst.font = _cambria()
        gst.alignment = Alignment(horizontal="center")
        gst.number_format = PERCENT_FMT

        qty_cell = ws.cell(row=row, column=4, value=qty)
        qty_cell.font = _cambria()
        qty_cell.alignment = Alignment(horizontal="center")

        rate_cell = ws.cell(row=row, column=5, value=rate)
        rate_cell.font = _cambria()
        rate_cell.alignment = Alignment(horizontal="center")
        rate_cell.number_format = RATE_FMT

        amt_cell = ws.cell(row=row, column=6, value=amount)
        amt_cell.font = _cambria()
        amt_cell.alignment = Alignment(horizontal="center")
        amt_cell.number_format = CURRENCY_FMT

        model_cell = ws.cell(row=row + 1, column=2, value=item.get("model", ""))
        model_cell.font = _cambria(bold=True)
        model_cell.alignment = Alignment(horizontal="left")

        for col in range(1, 7):
            _box(ws.cell(row=row, column=col), left=True, right=(col == 6))
            _box(ws.cell(row=row + 1, column=col), left=True, right=(col == 6))

    # ---- totals / terms / bank -------------------------------------------
    terms_row = last_item_row + 1
    total_row = terms_row + 2
    bank_head_row = terms_row + 3
    bank_body_row = terms_row + 4
    contact_row = terms_row + 7

    for col in range(1, 7):
        _box(ws.cell(row=terms_row, column=col), left=True, right=(col == 6))

    terms_head = ws.cell(row=terms_row, column=2, value="TERMS & CONDITIONS:-")
    terms_head.font = _cambria(bold=True, color=WHITE)
    terms_head.fill = RED_FILL
    _box(terms_head, left=True, top=True, bottom=True)

    gst_label = ws.cell(row=terms_row, column=5, value="GST AMT")
    gst_label.font = _cambria()
    gst_label.alignment = Alignment(horizontal="center", vertical="center")
    _box(gst_label, left=True, top=True)

    gst_value = ws.cell(row=terms_row, column=6, value=quote.get("gst_amount", 0))
    gst_value.font = _cambria()
    gst_value.alignment = Alignment(horizontal="center", vertical="center")
    gst_value.number_format = CURRENCY_FMT
    _box(gst_value, left=True, right=True, top=True, bottom=True)

    terms_block = _merge_write(
        ws, f"B{terms_row + 1}:B{terms_row + 6}", quote.get("terms", ""),
        font=_cambria(bold=True, color=WHITE),
        align=Alignment(horizontal="left", vertical="top", wrap_text=True),
        fill=BLACK_FILL,
    )
    _box(terms_block, left=True, bottom=True)

    total_label = ws.cell(row=total_row, column=5, value="G.TOTAL")
    total_label.font = _cambria()
    total_label.alignment = Alignment(horizontal="center", vertical="center")
    _box(total_label, left=True, top=True)

    total_value = ws.cell(row=total_row, column=6)
    # Live formula, exactly as the AppSheet template: the item amounts plus the
    # GST amount that sits in F<terms_row>.
    total_value.value = f"=SUM(F{FIRST_ITEM_ROW}:F{terms_row + 1})"
    total_value.font = _cambria()
    total_value.alignment = Alignment(horizontal="center", vertical="center")
    total_value.number_format = CURRENCY_FMT
    _box(total_value, left=True, right=True, top=True, bottom=True)

    bank_head = _merge_write(ws, f"C{bank_head_row}:E{bank_head_row}",
                             "OUR BANK DETAILS",
                             font=_cambria(bold=True, color=WHITE),
                             fill=BLUE_FILL)
    _box(bank_head, top=True, bottom=True)

    bank_body = _merge_write(
        ws, f"C{bank_body_row}:F{bank_body_row + 2}", company.get("bank_details", ""),
        font=_cambria(8, bold=True),
        align=Alignment(horizontal="left", vertical="top", wrap_text=True),
    )
    _box(bank_body, left=True, right=True, bottom=True)

    for r in range(terms_row + 1, contact_row):
        ws.row_dimensions[r].height = 12.75

    # ---- contact strip + footer ------------------------------------------
    contact = (company.get("contact_line") or "").strip()
    if not contact:
        tel = (company.get("tel") or "").strip()
        cell = (quote.get("salesperson_phone") or "").strip()
        contact = " ;    ".join(
            p for p in [f"CELL:- {cell}" if cell else "", f"TEL:- {tel}." if tel else ""] if p)
    _merge_write(ws, f"A{contact_row}:F{contact_row}", contact,
                 font=_cambria(bold=True), align=Alignment(horizontal="center"))
    ws.row_dimensions[contact_row].height = 14.25

    note_row = contact_row + 1
    for offset, note in enumerate(FOOTER_NOTES):
        cell = ws.cell(row=note_row + offset, column=1, value=note)
        cell.font = _cambria()
        cell.alignment = Alignment(horizontal="left")
        if offset in (0, 2, 3):
            ws.merge_cells(start_row=note_row + offset, start_column=1,
                           end_row=note_row + offset, end_column=2)

    sign_name = _merge_write(ws, f"D{note_row}:F{note_row}", company.get("name", ""),
                             font=_cambria(bold=True),
                             align=Alignment(horizontal="center", vertical="center"))
    ws.row_dimensions[note_row].height = 14.25
    for extra in (1, 2):
        ws.merge_cells(start_row=note_row + extra, start_column=4,
                       end_row=note_row + extra, end_column=6)

    decl_row = note_row + 4
    for offset, line in enumerate(GST_DECLARATION):
        cell = ws.cell(row=decl_row + offset, column=1, value=line)
        cell.font = _cambria()
        cell.alignment = Alignment(horizontal="left")
        end_col = 2 if offset == len(GST_DECLARATION) - 1 else 3
        ws.merge_cells(start_row=decl_row + offset, start_column=1,
                       end_row=decl_row + offset, end_column=end_col)

    sign_cell = ws.cell(row=decl_row + 3, column=4,
                        value="Proprietor / Authorised Signatory")
    sign_cell.font = _cambria()
    sign_cell.alignment = Alignment(horizontal="left")

    # ---- page setup -------------------------------------------------------
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.print_area = f"A1:{get_column_letter(6)}{decl_row + len(GST_DECLARATION)}"

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    return out_path


def xlsx_to_pdf(xlsx_path, out_dir=None):
    """Render the workbook to a one-page PDF using LibreOffice.

    Returns the PDF path, or None when LibreOffice is not installed.
    """
    out_dir = out_dir or os.path.dirname(os.path.abspath(xlsx_path))
    with tempfile.TemporaryDirectory() as profile:
        try:
            subprocess.run(
                ["soffice", "--headless", f"-env:UserInstallation=file://{profile}",
                 "--convert-to", "pdf", "--outdir", out_dir, xlsx_path],
                check=True, capture_output=True, timeout=180,
            )
        except (subprocess.SubprocessError, FileNotFoundError, OSError):
            return None
    pdf_path = os.path.join(
        out_dir, os.path.splitext(os.path.basename(xlsx_path))[0] + ".pdf")
    return pdf_path if os.path.exists(pdf_path) else None

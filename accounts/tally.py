"""Reads a Tally "Bills Payable" / "Bills Receivable" export.

The layout Tally writes:

    WELL WORTH INDUSTRIES -  (from 1-Apr-26)      <- firm, line 1
    ...address lines...
    Bills Receivable                               <- which report
    1-Apr-26 to 30-Oct-26
    Pending Bills
    Date | Ref. No. | Party's Name | Pending | Due on | Overdue
         |          |              | Amount  |        | by days
    22-Mar-24 | 400/WWI/23-24 | CHARU JEWELS | 32214 | 22-Mar-24 | 952
    ...
                                     | 4104873.09                   <- total

Tally saves these as .xls even though most are really .xlsx inside, so the
file is sniffed rather than trusted by its name.  A real old-style .xls is
read with xlrd, and a .csv works too.
"""

import csv
import io
import re
from datetime import date, datetime, timedelta

from .db import squeeze


class TallyFileError(ValueError):
    """The file is not a Bills Payable / Receivable report we can read."""


# ---------------------------------------------------------------- reading cells

def _rows_xlsx(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def _rows_xls(data):
    try:
        import xlrd
    except ImportError as exc:      # pragma: no cover - xlrd is in requirements
        raise TallyFileError("This is an old-style .xls file. Install xlrd, or "
                             "save it from Excel as .xlsx and upload again.") from exc
    book = xlrd.open_workbook(file_contents=data)
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                row.append(None)
            else:
                row.append(cell.value)
        rows.append(row)
    return rows


def _rows_csv(data):
    text = data.decode("utf-8-sig", errors="replace")
    return [[c if c != "" else None for c in row] for row in csv.reader(io.StringIO(text))]


def read_rows(data):
    if data[:2] == b"PK":                       # zip container: .xlsx
        return _rows_xlsx(data)
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":   # OLE2: real .xls
        return _rows_xls(data)
    head = data[:2000].lower()
    if b"<html" in head or b"<table" in head:
        raise TallyFileError("This file is a web page saved as .xls. Open it in "
                             "Excel, save it as .xlsx, and upload that.")
    return _rows_csv(data)


# ---------------------------------------------------------------- values

DATE_FORMATS = ("%d-%b-%y", "%d-%b-%Y", "%d-%m-%Y", "%d-%m-%y", "%d/%m/%Y",
                "%d/%m/%y", "%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d %H:%M:%S")


def parse_date(value):
    """A date cell -> 'YYYY-MM-DD', or '' when it is not a date."""
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        # an Excel serial number that came through a csv
        if 20000 < value < 80000:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).strftime("%Y-%m-%d")
        return ""
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return ""


_NUM = re.compile(r"-?\d+(?:\.\d+)?")


def parse_amount(value, kind=None):
    """32214 / '32,214.00 Dr' / '(1,200.00)' -> float, or None.

    A receivable is normally a debit and a payable a credit; the other side
    (an advance, an on-account payment) comes back negative."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    negative = text.startswith("(") and text.endswith(")")
    match = _NUM.search(text)
    if not match:
        return None
    amount = float(match.group())
    lowered = text.lower()
    if kind == "receivable" and lowered.endswith("cr"):
        negative = True
    if kind == "payable" and lowered.endswith("dr"):
        negative = True
    return -abs(amount) if negative else amount


def parse_int(value):
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        return 0


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# ---------------------------------------------------------------- the report

HEADER_NAMES = {
    "date": ("date",),
    "ref": ("ref", "refno", "billno", "invoiceno", "vchno"),
    "party": ("partysname", "party", "partyname", "name", "ledger"),
    "pending": ("pending", "pendingamount", "amount", "balance", "outstanding"),
    "due": ("dueon", "duedate"),
    "overdue": ("overdue", "overduebydays", "days"),
}


def _header_map(row):
    found = {}
    for idx, cell in enumerate(row):
        key = squeeze(cell).lower()
        if not key:
            continue
        for field, names in HEADER_NAMES.items():
            if field in found:
                continue
            if key in names or (field == "ref" and key.startswith("ref")):
                found[field] = idx
                break
    if {"date", "party", "pending"} <= set(found):
        return found
    return None


FILE_DATE = re.compile(r"(\d{1,2})[.\-_](\d{1,2})[.\-_](\d{2,4})")


def date_from_filename(name):
    """WWI_RECEIVABLE_23.10.26.xls -> 2026-10-23."""
    match = FILE_DATE.search(name or "")
    if not match:
        return ""
    d, m, y = (int(g) for g in match.groups())
    if y < 100:
        y += 2000
    try:
        return date(y, m, d).isoformat()
    except ValueError:
        return ""


def kind_from_text(text):
    t = squeeze(text).lower()
    if "receivable" in t or "debtor" in t:
        return "receivable"
    if "payable" in t or "creditor" in t:
        return "payable"
    return ""


def parse_report(data, file_name=""):
    """-> dict(firm_text, kind, bills=[...], total, report_date)."""
    rows = read_rows(data)
    if not rows:
        raise TallyFileError("The file is empty.")

    header_at, columns = None, None
    for i, row in enumerate(rows[:40]):
        columns = _header_map(row)
        if columns:
            header_at = i
            break
    if header_at is None:
        raise TallyFileError(
            "Could not find the Date / Ref. No. / Party's Name / Pending columns. "
            "Export the Bills Payable or Bills Receivable report from Tally and "
            "upload that file as it is.")

    top_lines = [_text(c) for row in rows[:header_at] for c in row if _text(c)]
    firm_text = top_lines[0] if top_lines else ""
    kind = ""
    for line in top_lines:
        kind = kind_from_text(line)
        if kind:
            break
    if not kind:
        kind = kind_from_text(file_name)

    bills = []
    total_row = None
    current_party = ""
    for row in rows[header_at + 1:]:
        def cell(field):
            idx = columns.get(field)
            return row[idx] if idx is not None and idx < len(row) else None

        bill_date = parse_date(cell("date"))
        party = _text(cell("party"))
        ref = _text(cell("ref"))
        amount = parse_amount(cell("pending"), kind)

        if not bill_date:
            if party and not ref and amount is None:
                current_party = party         # "group by party" layout
            elif not party and amount is not None and not ref:
                total_row = amount            # grand total at the bottom
            continue
        party = party or current_party
        if not party or amount is None:
            continue
        bills.append({
            "bill_date": bill_date,
            "ref_no": ref,
            "party": " ".join(party.split()),
            "pending": round(amount, 2),
            "due_date": parse_date(cell("due")) or bill_date,
            "overdue_days": parse_int(cell("overdue")),
        })

    if not bills:
        raise TallyFileError("The file has the right columns but no bills in it.")

    total = round(sum(b["pending"] for b in bills), 2)
    return {
        "firm_text": firm_text,
        "kind": kind,
        "bills": bills,
        "total": total,
        "total_row": total_row,
        "report_date": date_from_filename(file_name),
    }


def bill_key(bill):
    return "|".join((squeeze(bill["party"]), squeeze(bill["ref_no"]), bill["bill_date"]))
